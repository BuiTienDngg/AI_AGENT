import copy
from pathlib import Path
import tempfile
import unittest
from docx import Document

from game_seo.core import ROOT, load_json
from game_seo.export import export_all
from game_seo.qa import apply_patches, validate
from game_seo.translation import validate_translation
from test_pipeline import fixtures, translated_fixture


class CaptionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        self.article, self.sources, self.images, self.frames = fixtures(self.folder)
        self.cfg = load_json(ROOT / 'config.json')
        self.article['sections'][2]['subsections'] = [{'heading': 'Train Roof Routes',
            'paragraphs': [self.article['sections'][2]['paragraphs'].pop(0)]}]

    def tearDown(self):
        self.tmp.cleanup()

    def test_nearest_h3_then_h2_in_both_word_languages(self):
        vi = translated_fixture(self.article)
        vi['sections'][2]['subsections'][0]['heading'] = 'Đường trên nóc tàu'
        files = export_all(self.article, self.sources, self.images, self.frames, self.folder,
                           translation_vi=vi)
        doc_text = '\n'.join(p.text for p in Document(files['docx_en']).paragraphs)
        self.assertIn('Train Roof Routes', doc_text)
        self.assertNotIn('Gameplay at approximately', doc_text)
        en_captions = [p.text for p in Document(files['docx_en']).paragraphs if p.style.name == 'Caption']
        vi_captions = [p.text for p in Document(files['docx_vi']).paragraphs if p.style.name == 'Caption']
        self.assertTrue(any(p.endswith('Train Roof Routes') for p in en_captions))
        self.assertTrue(any(p.endswith('Gameplay') for p in en_captions))
        self.assertTrue(any(p.endswith('Đường trên nóc tàu') for p in vi_captions))
        self.assertTrue(any(p.endswith('Lối chơi') for p in vi_captions))
        self.assertTrue(vi_captions[0].startswith('Hình 1. Ý tưởng và cảm hứng'))

    def test_specific_caption_overrides_heading_and_keeps_credits(self):
        paragraph = self.article['sections'][0]['paragraphs'][0]
        paragraph['image_captions'] = [{'image_id': 'IMG001', 'caption': 'Routes <above> & below the city'}]
        vi = translated_fixture(self.article)
        vi['sections'][0]['paragraphs'][0]['image_captions'][0]['caption'] = 'Những tuyến đường trong thành phố'
        files = export_all(self.article, self.sources, self.images, self.frames, self.folder, translation_vi=vi)
        en_doc = Document(files['docx_en'])
        en_captions = [p.text for p in en_doc.paragraphs if p.style.name == 'Caption']
        self.assertTrue(en_captions[0].endswith(paragraph['image_captions'][0]['caption']))
        credits = '\n'.join(p.text for p in en_doc.paragraphs if p.style.name == 'Image Credit')
        self.assertIn('Fixture Photographer', credits)
        self.assertIn('CC BY-SA', credits)
        vi_captions = [p.text for p in Document(files['docx_vi']).paragraphs if p.style.name == 'Caption']
        self.assertTrue(vi_captions[0].endswith('Những tuyến đường trong thành phố'))
        self.assertTrue(validate_translation(self.article, vi)['pass'])

    def test_removed_or_reassigned_translated_caption_is_blocked(self):
        paragraph = self.article['sections'][0]['paragraphs'][0]
        paragraph['image_captions'] = [{'image_id': 'IMG001', 'caption': 'Urban routes'}]
        for broken in ([], [{'image_id': 'IMG002', 'caption': 'Tuyến đường'}]):
            with self.subTest(broken=broken):
                vi = translated_fixture(self.article)
                vi['sections'][0]['paragraphs'][0]['image_captions'] = broken
                self.assertIn('CAPTION_SEQUENCE', [x['code'] for x in validate_translation(self.article, vi)['issues']])

    def test_empty_translated_caption_is_blocked(self):
        self.article['sections'][0]['paragraphs'][0]['image_captions'] = [{'image_id': 'IMG001', 'caption': 'Urban routes'}]
        vi = translated_fixture(self.article)
        vi['sections'][0]['paragraphs'][0]['image_captions'][0]['caption'] = ''
        self.assertIn('EMPTY_IMAGE_CAPTION', [x['code'] for x in validate_translation(self.article, vi)['issues']])

    def test_invalid_caption_ids_and_internal_labels_fail_qa(self):
        for item, expected in [({'image_id': 'IMG002', 'caption': 'Urban routes'}, 'IMAGE_CAPTION_SEQUENCE'),
                               ({'image_id': 'IMG001', 'caption': 'FRAME001 at 12s'}, 'INVALID_IMAGE_CAPTION')]:
            with self.subTest(item=item):
                self.article['sections'][0]['paragraphs'][0]['image_captions'] = [item]
                result = validate(self.article, self.sources, self.images, self.frames, self.cfg)
                self.assertIn(expected, [x['code'] for x in result['issues']])

    def test_patch_replaces_caption_and_drops_stale_legacy_caption(self):
        self.article['sections'][0]['paragraphs'][0]['image_captions'] = [{'image_id': 'IMG001', 'caption': 'Old caption'}]
        edit = {'path': 'sections[0].paragraphs[0]', 'text': 'New paragraph [[SRC001]]', 'image_ids': ['IMG002'],
                'image_captions': [{'image_id': 'IMG002', 'caption': 'New caption'}]}
        updated, _ = apply_patches(self.article, {'edits': [edit]})
        self.assertEqual(updated['sections'][0]['paragraphs'][0]['image_captions'], edit['image_captions'])
        del edit['image_captions']
        legacy, _ = apply_patches(self.article, {'edits': [edit]})
        self.assertNotIn('image_captions', legacy['sections'][0]['paragraphs'][0])
