from __future__ import annotations
import html
import re
from pathlib import Path
from docx import Document
from docx.image.exceptions import UnrecognizedImageError
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .qa import TOKEN_PATTERN, walk_paragraphs, verify_references
from .core import slugify, save_json
from .translation import validate_translation


def numbered_text(text, numbers):
    return TOKEN_PATTERN.sub(lambda m: f'[{numbers[m.group(1)]}]', text)


def image_caption(paragraph, image_id, heading):
    """Prefer an editorial caption; legacy drafts use their nearest H3/H2."""
    for item in paragraph.get('image_captions', []):
        if item.get('image_id') == image_id and str(item.get('caption', '')).strip():
            return str(item['caption']).strip()
    return heading


def add_citation_runs(paragraph, text, numbers):
    position = 0
    for match in TOKEN_PATTERN.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        ref_id = match.group(1)
        citation = paragraph.add_run(f'[{numbers[ref_id]}]')
        citation.bold = True
        citation.italic = True
        citation.font.size = Pt(8)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def figure(doc, img, label, kind, lang='en-US'):
    path = Path(img.get('path', ''))
    if not path.exists():
        raise RuntimeError(f'Expected local figure missing: {path}')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    from PIL import Image
    with Image.open(path) as source_image:
        w, h = source_image.size
    scaling = min(5.8 / max(w, 1), 5.85 / max(h, 1))
    try:
        run.add_picture(str(path), width=Inches(w * scaling), height=Inches(h * scaling))
    except UnrecognizedImageError:
        # Some valid Commons JPEGs lack the JFIF/Exif signature python-docx expects.
        # Normalize only the embedded copy; keep original assets and attribution.
        from io import BytesIO
        normalized = BytesIO()
        with Image.open(path) as source_image:
            source_image.convert('RGBA').save(normalized, format='PNG')
        normalized.seek(0)
        run.add_picture(normalized, width=Inches(w * scaling), height=Inches(h * scaling))
    caption = doc.add_paragraph(style='Caption')
    caption.paragraph_format.keep_together = True
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.add_run(label)
    if kind == 'Wikimedia Commons' and img.get('author'):
        credit = doc.add_paragraph(style='Image Credit')
        credit.alignment = WD_ALIGN_PARAGRAPH.CENTER
        photo_label = 'Ảnh' if lang == 'vi-VN' else 'Photo'
        credit.add_run(f"{photo_label}: {img['author']} | {img['license']}")


def add_paras(doc, paragraphs, numbers, images_by_id, frames_by_id, fig_index, heading='', lang='en-US'):
    for para in paragraphs:
        text = str(para.get('text', '')).strip()
        if text:
            p = doc.add_paragraph(style='Normal')
            add_citation_runs(p, text, numbers)
        for ref in para.get('image_ids', []):
            fig_index += 1
            figure_label = 'Hình' if lang == 'vi-VN' else 'Figure'
            label = f'{figure_label} {fig_index}. {image_caption(para, ref, heading)}'
            if ref in images_by_id:
                img = images_by_id[ref]
                figure(doc, img, label, 'Wikimedia Commons', lang=lang)
            elif ref in frames_by_id:
                frame = frames_by_id[ref]
                figure(doc, frame, label, 'Gameplay capture', lang=lang)
            else:
                raise RuntimeError(f'Unregistered image ID {ref}')
    return fig_index


def export_docx(article, source_registry, image_registry, frames, out_path: Path, lang='en-US'):
    if lang not in ('en-US', 'vi-VN'):
        raise ValueError(f'Unsupported Word language: {lang}')
    source_by_id = {x['id']: x for x in source_registry}
    image_by_id = {x['id']: x for x in image_registry}
    frame_by_id = {x['frame_id']: x for x in frames}
    numbers = verify_references(source_registry, article)
    doc = Document()
    doc.core_properties.language = lang
    defaults = doc.styles.element.find(qn('w:docDefaults'))
    if defaults is not None:
        default_lang = defaults.find('.//' + qn('w:lang'))
        if default_lang is not None:
            default_lang.set(qn('w:val'), lang)
    section = doc.sections[0]
    section.top_margin = Inches(.78)
    section.bottom_margin = Inches(.74)
    section.left_margin = Inches(.86)
    section.right_margin = Inches(.86)
    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(10.5)
    proofing = OxmlElement('w:lang')
    proofing.set(qn('w:val'), lang)
    normal.element.get_or_add_rPr().append(proofing)
    normal.paragraph_format.space_after = Pt(7)
    for style_name, size in [('Title', 20), ('Heading 1', 15), ('Heading 2', 12)]:
        st = doc.styles[style_name]
        st.font.name = 'Aptos Display' if style_name == 'Title' else 'Aptos'
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(29, 48, 68)
    if 'Image Credit' not in doc.styles:
        credit = doc.styles.add_style('Image Credit', WD_STYLE_TYPE.PARAGRAPH)
        credit.base_style = normal
        credit.font.name = 'Aptos'
        credit.font.size = Pt(8)
        credit.font.italic = True
        credit.font.color.rgb = RGBColor(82, 91, 103)
    doc.add_paragraph(article['title'], style='Title')
    fig_index = 0
    for s in article['sections']:
        doc.add_heading(s['heading'], 1)
        fig_index = add_paras(doc, s.get('paragraphs', []), numbers, image_by_id, frame_by_id, fig_index, s['heading'], lang=lang)
        for sub in s.get('subsections', []):
            doc.add_heading(sub['heading'], 2)
            fig_index = add_paras(doc, sub.get('paragraphs', []), numbers, image_by_id, frame_by_id, fig_index, sub['heading'], lang=lang)
    doc.add_heading('Tài liệu tham khảo' if lang == 'vi-VN' else 'References', 1)
    for sid, n in sorted(numbers.items(), key=lambda it: it[1]):
        source = source_by_id[sid]
        doc.add_paragraph(f"[{n}] {source['title']}. {source['url']}")
    commons = [image_by_id[k] for k in dict.fromkeys(ref for _, _, p in walk_paragraphs(article) for ref in p.get('image_ids', [])) if k in image_by_id]
    if commons:
        doc.add_heading('Nguồn ảnh' if lang == 'vi-VN' else 'Image Credits', 1)
        for img in commons:
            doc.add_paragraph(f"{img['title']} — {img['author']} — {img['license']}. {img['page_url']}" +
                              (f" {'Giấy phép' if lang == 'vi-VN' else 'License'}: {img['license_url']}" if img.get('license_url') else ''))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Slope3 | Bản thảo bài game | vi-VN' if lang == 'vi-VN' else 'Slope3 | Game Editorial Draft | en-US')
    out_path.parent.mkdir(exist_ok=True, parents=True)
    doc.save(out_path)
    return out_path


def markdown_body(article, sources, images, frames):
    numbers = verify_references(sources, article)
    idx_sources = {s['id']: s for s in sources}
    idx_images = {x['id']: x for x in images}
    idx_frames = {x['frame_id']: x for x in frames}
    lines = [f"# {article['title']}", '', f"*{article.get('meta_description', '')}*", '']
    for sec in article['sections']:
        lines.extend([f"## {sec['heading']}", ''])
        sections = [(None, sec.get('paragraphs', []))] + [(sub['heading'], sub.get('paragraphs', [])) for sub in sec.get('subsections', [])]
        for heading, paragraphs in sections:
            if heading:
                lines.extend([f'### {heading}', ''])
            for para in paragraphs:
                lines.extend([numbered_text(para.get('text', ''), numbers), ''])
                for imgid in para.get('image_ids', []):
                    caption = image_caption(para, imgid, heading or sec['heading'])
                    if imgid in idx_images:
                        img = idx_images[imgid]
                        lines.extend([f"![{caption}]({Path(img['path']).as_posix()})", '', f'*{caption}*', '',
                                      f"*Image: {img['author']} / {img['license']} — {img['page_url']}*", ''])
                    elif imgid in idx_frames:
                        frame = idx_frames[imgid]
                        lines.extend([f"![{caption}]({Path(frame['path']).as_posix()})", '', f'*{caption}*', ''])
    lines.extend(['## References', ''])
    for sid, n in sorted(numbers.items(), key=lambda it: it[1]):
        src = idx_sources[sid]
        lines.append(f"{n}. {src['title']} — {src['url']}")
    lines.extend(['', '## Image Credits', ''])
    seen = set()
    for _, _, p in walk_paragraphs(article):
        for imgid in p.get('image_ids', []):
            if imgid in idx_images and imgid not in seen:
                seen.add(imgid)
                img = idx_images[imgid]
                lines.append(f"- {img['title']} — {img['author']} — {img['license']} — {img['page_url']}")
    return '\n'.join(lines) + '\n'


def export_html(article, sources, images, frames, outfile, lang="en-US"):
    if lang not in ("en-US", "vi"):
        raise ValueError(f"Unsupported HTML language: {lang}")
    refs_heading = "Tài liệu tham khảo" if lang == "vi" else "References"
    credits_heading = "Nguồn ảnh" if lang == "vi" else "Image Credits"
    gameplay_label = "Ảnh từ trò chơi" if lang == "vi" else "Gameplay capture"
    img_source_label = "Nguồn" if lang == "vi" else "Source"
    numbers = verify_references(sources, article)
    src_map = {s['id']: s for s in sources}
    image_map = {x['id']: x for x in images}
    frame_map = {f['frame_id']: f for f in frames}
    def format_paragraph(s):
        result = []
        pos = 0
        for m in TOKEN_PATTERN.finditer(s):
            result.append(html.escape(s[pos:m.start()]))
            n = numbers[m.group(1)]
            result.append(f'<a href="#ref-{n}" class="cite">[{n}]</a>')
            pos = m.end()
        result.append(html.escape(s[pos:]))
        return ''.join(result)
    lines = ['<!doctype html>', f'<html lang="{lang}"><head><meta charset="utf-8">',
             '<meta name="viewport" content="width=device-width, initial-scale=1">',
             f'<meta name="description" content="{html.escape(article.get("meta_description", ""), quote=True)}">',
             f'<title>{html.escape(article.get("meta_title") or article["title"])}</title>',
             '<style>body{max-width:800px;margin:45px auto;padding:0 22px;font:17px/1.65 system-ui,Arial,sans-serif;color:#202733}img{display:block;max-width:100%;height:auto;margin:16px auto}figcaption{color:#555;font-size:.8em}h2{margin-top:38px} .cite{font-size:.82em;font-weight:700}li{overflow-wrap:anywhere}</style></head><body>',
             f'<h1>{html.escape(article["title"])}</h1>']
    def emit_paragraphs(paragraphs, heading):
        for para in paragraphs:
            lines.append('<p>' + format_paragraph(para.get('text', '')) + '</p>')
            for ref in para.get('image_ids', []):
                meta = image_map.get(ref) or frame_map.get(ref)
                if not meta:
                    continue
                # file:// is deliberately not used; local export includes a relative assets path.
                local = Path(meta['path']).resolve()
                relative = f'assets/{local.name}'
                credit = (f"{meta['author']} / {meta['license']}" if ref in image_map else gameplay_label)
                caption = image_caption(para, ref, heading)
                lines.append(f'<figure><img src="{html.escape(relative, quote=True)}" alt="{html.escape(caption, quote=True)}"><figcaption><span class="image-caption">{html.escape(caption)}</span><br><small class="image-credit">{html.escape(credit)}</small></figcaption></figure>')
    for sec in article['sections']:
        lines.append('<h2>' + html.escape(sec['heading']) + '</h2>')
        emit_paragraphs(sec.get('paragraphs', []), sec['heading'])
        for sub in sec.get('subsections', []):
            lines.append('<h3>' + html.escape(sub['heading']) + '</h3>')
            emit_paragraphs(sub.get('paragraphs', []), sub['heading'])
    lines.append(f'<h2>{refs_heading}</h2><ol>')
    for sid, n in sorted(numbers.items(), key=lambda x: x[1]):
        s = src_map[sid]
        lines.append(f'<li id="ref-{n}">{html.escape(s["title"])} — <a href="{html.escape(s["url"], quote=True)}">{html.escape(s["url"])}</a></li>')
    lines.append(f'</ol><h2>{credits_heading}</h2><ul>')
    used_images = list(dict.fromkeys(ref for _, _, p in walk_paragraphs(article) for ref in p.get('image_ids', []) if ref in image_map))
    for ref in used_images:
        im = image_map[ref]
        lines.append(f'<li>{html.escape(im["title"])} — {html.escape(im["author"])} — {html.escape(im["license"])} — <a href="{html.escape(im["page_url"], quote=True)}">{img_source_label}</a></li>')
    lines.extend(['</ul></body></html>'])
    outfile.parent.mkdir(parents=True, exist_ok=True)
    outfile.write_text('\n'.join(lines), encoding='utf-8')
    import shutil
    assets = outfile.parent / 'assets'
    assets.mkdir(exist_ok=True)
    all_ids = list(dict.fromkeys(ref for _, _, p in walk_paragraphs(article) for ref in p.get('image_ids', [])))
    for ref in all_ids:
        m = image_map.get(ref) or frame_map.get(ref)
        if m:
            shutil.copy2(m['path'], assets / Path(m['path']).name)
    return outfile


def export_all(article, sources, images, frame_manifest, workdir, translation_vi=None, *, editorial_v2=False):
    """Export exactly two Word documents from the approved, aligned EN/VI text."""
    if translation_vi is None:
        raise ValueError('translation_vi is required to create both English and Vietnamese Word documents')
    alignment = validate_translation(article, translation_vi,localize_title=editorial_v2)
    if not alignment['pass']:
        raise ValueError(f'VI translation failed structural QA: {alignment["issues"][:5]}')
    out = Path(workdir) / 'output'
    out.mkdir(parents=True, exist_ok=True)
    base = slugify(article['title']) + '_slope3'
    docx_path = out / f'{base}_en_US.docx'
    docx_vi = out / f'{base}_vi_VN.docx'
    # Files are only generated after translation consistency validation succeeds.
    export_docx(article, sources, images, frame_manifest['frames'], docx_path)
    export_docx(translation_vi, sources, images, frame_manifest['frames'], docx_vi, lang='vi-VN')
    files = {'docx_en': str(docx_path), 'docx_vi': str(docx_vi)}
    from .result_bundle import build_result
    build_result(article, translation_vi, images, frame_manifest, files, workdir)
    return files
