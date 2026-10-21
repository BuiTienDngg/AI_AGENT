from pathlib import Path
import copy
import json
from unittest.mock import patch
from types import SimpleNamespace
import tempfile
import unittest
from PIL import Image
from game_seo.core import ROOT, load_json, save_json
from game_seo.qa import validate, apply_patches, verify_references
from game_seo.export import export_all
from game_seo.translation import validate_translation, HEADINGS_VI
from game_seo.usage import extract_usage, record_codex_call
from game_seo.runner import codex_call
from game_seo.research import safe_license, academic_abstract
from game_seo.runner import filter_notes, build_repair_context


def fixtures(folder):
    folder = Path(folder)
    folder.mkdir(exist_ok=True, parents=True)
    frames = []
    images = []
    for i in range(1, 5):
        path = folder / f'FRAME{i:03}.jpg'
        Image.new('RGB', (720, 405), (80+i*20, 120, 160)).save(path)
        frames.append({'frame_id': f'FRAME{i:03}', 'path': str(path), 'timestamp_seconds_approx': i * 12})
    for i in range(1, 3):
        path = folder / f'IMG{i:03}.jpg'
        Image.new('RGB', (420, 720) if i == 2 else (720, 420), (90, 110+i*35, 130)).save(path)
        images.append({'id': f'IMG{i:03}', 'title': 'Metro architecture photo', 'path': str(path),
                       'author': f'Fixture Photographer {i}', 'license': 'CC BY-SA 3.0',
                       'page_url': f'https://commons.wikimedia.org/wiki/File:example_photo_{i}.jpg',
                       'license_url': 'https://creativecommons.org/licenses/by-sa/3.0/'})
    sources = [
        {'id': f'SRC{i:03}', 'title': f'Fixture source {i}', 'url': f'https://example.org/valid-source-{i}',
         'excerpt': 'This example source describes game design and motivation with explanatory background sufficient for testing only. ' * 3,
         'kind': 'academic' if i == 2 else 'encyclopedia', 'verified_metadata': True, 'evidence_quality': 'abstract' if i == 2 else 'overview'}
        for i in range(1, 5)
    ]
    intro = [('Urban transit environments shape the feeling of movement and wayfinding, setting a suitable visual analogy for this runner. [[SRC001]]', ['IMG001']),
             ('Inspiration in this article is interpretive: it connects public infrastructure to the interactive obstacle course, not the developer’s confirmed original motivation. [[SRC003]]', ['IMG002'])]
    sections = []
    items = [
        ('Concept and Inspiration', intro),
        ('Story', [('The running character finds another track, another turn, and another tunnel. This creative vignette depicts a possible journey through the scene, not established canon.', [])]),
        ('Gameplay', [('The avatar moves forward through railcars and collectibles; short reactions matter. The interaction loop is continuous, readable, and varied. [[SRC004]]', ['FRAME001', 'FRAME002']),
                      ('Obstacles narrow the available lanes, and coins draw attention along a path through the urban backdrop.', ['FRAME003'])]),
        ('Visual Style', [('Contrasting environmental colors separate useful objects and visual background detail; clean silhouettes help the player read the path ahead. [[SRC003]]', ['FRAME004'])]),
        ('Standout Features', [('The game combines simple actions with a sense of escalating pressure as urban obstacles appear. Related design discussions emphasize readable challenges. [[SRC004]]', [])]),
        ('Controls', [('The exact keyboard and touchscreen mappings require confirmation; this article instead describes the need to switch lanes and react to obstacles, without inventing inputs. [[SRC002]]', [])]),
    ]
    for name, paras in items:
        sections.append({'heading': name, 'paragraphs': [{'text': t, 'image_ids': imgs} for t, imgs in paras], 'subsections': []})
    # The synthetic fixture intentionally satisfies the minimum article-length gate.
    for section in sections:
        for paragraph in section['paragraphs']:
            paragraph['text'] += (' The scenario is a controlled editorial test showing how the layout preserves separation among gameplay observation, cultural interpretation, source evidence, and publisher claims.')
    article = {'title': 'Fixture Game: An Urban Endless Runner', 'meta_title': 'Fixture Game | Play the Runner', 'meta_description': 'Discover the obstacle rhythm and immersive urban design behind a fast-paced runner with a clear visual identity and a quick restart loop.', 'sections': sections}
    return article, sources, images, {'frames': frames}


def translated_fixture(article):
    result = copy.deepcopy(article)
    result['meta_title'] = 'Trò chơi mô phỏng chạy không ngừng'
    result['meta_description'] = 'Bản dịch tiếng Việt của nội dung giới thiệu game, giữ nguyên hình minh họa và trích dẫn nguồn.'
    for section in result['sections']:
        section['heading'] = HEADINGS_VI[section['heading']]
        for para in section['paragraphs']:
            tokens = __import__('re').findall(r'\[\[SRC\d{3}\]\]', para['text'])
            para['text'] = 'Bài viết tiếng Việt được dịch từ bản tiếng Anh đã được phê duyệt. ' + ' '.join(tokens)
    return result


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        self.article, self.sources, self.images, self.frames = fixtures(self.folder)
        self.cfg = load_json(ROOT / 'config.json')
    def tearDown(self):
        self.tmp.cleanup()
    def test_valid_citations_and_images_pass(self):
        r = validate(self.article, self.sources, self.images, self.frames, self.cfg)
        self.assertTrue(r['pass'], r['issues'])
        self.assertEqual(r['summary']['unique_sources'], 4)
        self.assertEqual(r['summary']['source_mentions'], 6)
    def test_unsourced_psychology_fails(self):
        self.article['sections'][0]['paragraphs'][0]['text'] = 'Research shows that gaming is fun.'
        r = validate(self.article, self.sources, self.images, self.frames, self.cfg)
        self.assertFalse(r['pass'])
        self.assertIn('ACADEMIC_CLAIM_UNSOURCED', [e['code'] for e in r['issues']])
    def test_fabricated_id_blocked(self):
        self.article['sections'][0]['paragraphs'][0]['text'] += ' [[SRC999]]'
        self.assertFalse(validate(self.article, self.sources, self.images, self.frames, self.cfg)['pass'])
    def test_intro_missing_photo_fails(self):
        self.article['sections'][0]['paragraphs'][1]['image_ids'] = []
        self.assertIn('INTRO_MISSING_COMMONS_IMAGES', [e['code'] for e in validate(self.article, self.sources, self.images, self.frames, self.cfg)['issues']])
    def test_patches_only_touch_target(self):
        old = self.article['sections'][1]['paragraphs'][0]['text']
        a, changed = apply_patches(self.article, {'edits': [{'path': 'sections[0].paragraphs[1]', 'text': 'A safe replacement [[SRC003]]', 'image_ids': ['IMG002']}]})
        self.assertEqual(changed, ['sections[0].paragraphs[1]'])
        self.assertEqual(a['sections'][1]['paragraphs'][0]['text'], old)
        self.assertEqual(self.article['sections'][0]['paragraphs'][1]['text'].startswith('Inspiration'), True)
    def test_unverified_source_notes_rejected(self):
        notes = {'context': [{'claim': 'invented', 'source_id': 'SRC999', 'supporting_excerpt': 'not real'}], 'player_psychology': [], 'genre_and_competitors': [], 'unsupported_questions': []}
        clean, rejected = filter_notes(notes, self.sources)
        self.assertEqual(clean['context'], [])
        self.assertEqual(len(rejected), 1)
    def test_license(self):
        self.assertIsNotNone(safe_license({'LicenseShortName': {'value': 'CC BY-SA 4.0'}, 'LicenseUrl': {'value': 'https://creativecommons.org/licenses/by-sa/4.0/'}, 'Artist': {'value': 'Some Author'}}))
        self.assertIsNone(safe_license({'LicenseShortName': {'value': 'CC BY-NC-ND 4.0'}, 'LicenseUrl': {'value': 'https://creativecommons.org/licenses/by-nc-nd/4.0/'}}))
    def test_abstract(self):
        self.assertEqual(academic_abstract({'Enjoyable': [1], 'games': [0], 'matter.': [2]}), 'games Enjoyable matter.')
    def test_only_two_localized_word_documents_with_matching_citations_and_media(self):
        report = validate(self.article, self.sources, self.images, self.frames, self.cfg)
        self.assertTrue(report['pass'])
        files = export_all(self.article, self.sources, self.images, self.frames, self.folder,
                           translation_vi=translated_fixture(self.article))
        self.assertEqual(set(files), {'docx_en', 'docx_vi'})
        self.assertEqual({p.name for p in (self.folder / 'output').iterdir()},
                         {Path(path).name for path in files.values()})
        from docx import Document
        doc = Document(files['docx_en'])
        all_text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('References', all_text)
        self.assertIn('Image Credits', all_text)
        self.assertNotIn('[[SRC', all_text)
        self.assertEqual(len(doc.inline_shapes), 6)
        # All in-body citations are styled bold, italic and 8 pt as requested.
        citations = [r for p in doc.paragraphs for r in p.runs if r.text.startswith('[') and r.text.endswith(']')]
        self.assertGreaterEqual(len(citations), 6)
        self.assertTrue(all(r.bold and r.italic and abs(r.font.size.pt-8)<.1 for r in citations))
        vi_doc = Document(files['docx_vi'])
        vi_text = '\n'.join(p.text for p in vi_doc.paragraphs)
        self.assertIn('Tài liệu tham khảo', vi_text)
        self.assertIn('Nguồn ảnh', vi_text)
        self.assertIn('Ý tưởng và cảm hứng', vi_text)
        self.assertIn('Hình 1.', vi_text)
        self.assertNotIn('[[SRC', vi_text)
        self.assertEqual(len(vi_doc.inline_shapes), 6)
        self.assertEqual(doc.core_properties.language, 'en-US')
        self.assertEqual(vi_doc.core_properties.language, 'vi-VN')
        self.assertIn('vi-VN', vi_doc.sections[0].footer.paragraphs[0].text)
        vi_citations = [r for p in vi_doc.paragraphs for r in p.runs if r.text.startswith('[') and r.text.endswith(']')]
        self.assertEqual([r.text for r in citations], [r.text for r in vi_citations])
        self.assertTrue(all(r.bold and r.italic and abs(r.font.size.pt-8)<.1 for r in vi_citations))
        for image in self.images:
            self.assertIn(image['author'], vi_text)
            self.assertIn(image['license_url'], vi_text)

    def test_bad_translation_blocks_both_word_outputs(self):
        vi = translated_fixture(self.article)
        vi['sections'][0]['paragraphs'][0]['image_ids'] = []
        with self.assertRaisesRegex(ValueError, 'translation failed'):
            export_all(self.article, self.sources, self.images, self.frames, self.folder, translation_vi=vi)
        self.assertFalse((self.folder / 'output').exists())

    def test_export_cli_reuses_translation_without_calling_codex_and_updates_manifest(self):
        import run
        save_json(self.folder / 'article.json', self.article)
        save_json(self.folder / 'sources.json', {'sources': self.sources})
        save_json(self.folder / 'images.json', {'images': self.images})
        save_json(self.folder / 'frames.json', self.frames)
        save_json(self.folder / 'critique.json', {'pass_editorial': True, 'issues': []})
        save_json(self.folder / 'translation_vi.json', translated_fixture(self.article))
        save_json(self.folder / 'translation_provenance.json', {'en_article_sha256': run.article_hash(self.article)})
        with patch('sys.argv', ['run.py', 'export', '--run-dir', str(self.folder)]), \
             patch('run.codex_call', side_effect=AssertionError('Export must not spend Codex quota for unchanged text')):
            self.assertEqual(run.cli(), 0)
        files = load_json(self.folder / 'export_manifest.json')['files']
        self.assertEqual(set(files), {'docx_en', 'docx_vi'})
        self.assertTrue(all(Path(path).is_file() for path in files.values()))

    def test_translation_validation_passes(self):
        vi = translated_fixture(self.article)
        self.assertTrue(validate_translation(self.article, vi)['pass'])

    def test_translation_citation_removed_is_blocked(self):
        vi = translated_fixture(self.article)
        vi['sections'][0]['paragraphs'][0]['text'] = 'Bị mất nguồn.'
        report = validate_translation(self.article, vi)
        self.assertFalse(report['pass'])
        self.assertIn('CITATION_SEQUENCE', [x['code'] for x in report['issues']])

    def test_translation_figure_order_modified_is_blocked(self):
        vi = translated_fixture(self.article)
        vi['sections'][0]['paragraphs'][0]['image_ids'] = []
        report = validate_translation(self.article, vi)
        self.assertIn('IMAGE_SEQUENCE', [x['code'] for x in report['issues']])

    def test_translation_heading_modified_is_blocked(self):
        vi = translated_fixture(self.article)
        vi['sections'][0]['heading'] = 'Lịch sử phát triển'
        report = validate_translation(self.article, vi)
        self.assertIn('VI_HEADING', [x['code'] for x in report['issues']])

    def test_cli_token_usage_from_turn_completed(self):
        stream = '\n'.join(json.dumps(x) for x in [
            {'type': 'thread.started', 'thread_id': 'abc'},
            {'type': 'turn.completed', 'usage': {'input_tokens': 1200, 'cached_input_tokens': 200, 'output_tokens': 250}},
            {'type': 'turn.completed', 'usage': {'input_tokens': 100, 'cached_input_tokens': 0, 'output_tokens': 50}},
        ])
        usage = extract_usage(stream)
        self.assertEqual((usage['input_tokens'], usage['output_tokens']), (1300, 300))
        self.assertEqual(usage['cached_input_tokens'], 200)

    def test_cli_token_usage_cumulative_fallback_no_double_count(self):
        stream = '\n'.join(json.dumps(x) for x in [
            {'type': 'thread.token_usage.updated', 'token_usage': {'total_token_usage': {'input_tokens': 100, 'cached_input_tokens': 10, 'output_tokens': 30}}},
            {'type': 'thread.token_usage.updated', 'token_usage': {'total_token_usage': {'input_tokens': 250, 'cached_input_tokens': 20, 'output_tokens': 75}}},
        ])
        usage = extract_usage(stream)
        self.assertEqual(usage['input_tokens'], 250)
        self.assertEqual(usage['output_tokens'], 75)

    def test_token_report_mark_missing_usage_without_fabricating(self):
        record_codex_call(self.folder, 'draft', 0,
            {'input_tokens': 2500, 'cached_input_tokens': 500, 'output_tokens': 1000, 'source': 'codex_cli_turn_completed'},
            7.2, 'Generate an article', '{}', 'codex_cli_default')
        record_codex_call(self.folder, 'translate', 0, None, 8.1, 'Translate', '{}', '')
        data = json.loads((self.folder / 'token_usage.json').read_text())
        self.assertEqual(data['measurement_status'], 'partial')
        self.assertEqual(data['measured_totals']['total_tokens'], 3500)
        self.assertEqual(data['calls_without_usage'], 1)
        self.assertEqual(data['measured_totals']['cached_input_tokens'], 500)

    def test_codex_cli_stage_json_events_are_measured(self):
        calls = []
        def fake_call(command, **kwargs):
            calls.append(command)
            result_file = Path(command[command.index('--output-last-message') + 1])
            result_file.write_text('{}', encoding='utf-8')
            events = json.dumps({'type': 'turn.completed', 'usage': {
                'input_tokens': 1700, 'cached_input_tokens': 400, 'output_tokens': 120}})
            return SimpleNamespace(returncode=0, stdout=events + '\n', stderr='')
        with patch('game_seo.runner.shutil.which', return_value='/fake/codex'), \
             patch('game_seo.runner.subprocess.run', side_effect=fake_call):
            self.assertEqual(codex_call('notes', self.folder, self.cfg), {})
            self.assertEqual(codex_call('notes', self.folder, self.cfg), {})
        self.assertTrue(all('--json' in c for c in calls))
        self.assertEqual(len(list((self.folder / 'logs').glob('codex_notes_*.jsonl'))), 2)
        report = json.loads((self.folder / 'token_usage.json').read_text())
        self.assertEqual(report['measured_totals']['total_tokens'], 2 * (1700 + 120))
        self.assertEqual(report['calls_total'], 2)


if __name__ == '__main__':
    unittest.main()
